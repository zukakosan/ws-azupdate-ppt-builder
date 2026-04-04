from pptx import Presentation

prs = Presentation('template/template.pptx')
print('=== Slide Masters ===')
for i, master in enumerate(prs.slide_masters):
    print(f'Master {i}')

print()
print('=== Slide Layouts ===')
for i, layout in enumerate(prs.slide_layouts):
    print(f'Layout {i}: "{layout.name}"')
    for j, ph in enumerate(layout.placeholders):
        print(f'  Placeholder {j}: idx={ph.placeholder_format.idx}, type={ph.placeholder_format.type}, name="{ph.name}", size=({ph.left},{ph.top},{ph.width},{ph.height})')

print()
print('=== Existing Slides ===')
print(f'Number of slides: {len(prs.slides)}')
for i, slide in enumerate(prs.slides):
    print(f'Slide {i}: layout="{slide.slide_layout.name}"')
    for sh in slide.shapes:
        print(f'  Shape: name="{sh.name}", type={sh.shape_type}, pos=({sh.left},{sh.top},{sh.width},{sh.height})')
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                txt = p.text[:80]
                print(f'    Text: "{txt}"')

print()
print(f'Width: {prs.slide_width}, Height: {prs.slide_height}')
print(f'Width (inches): {prs.slide_width/914400}, Height (inches): {prs.slide_height/914400}')
