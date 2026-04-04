"""
Azure Updates PPTX Generator
テンプレート template/template.pptx を使用し、output/updates_ja.json から
PowerPoint スライドを生成する。
"""

import json
import re
from collections import Counter, OrderedDict
from datetime import datetime

from pptx import Presentation
from pptx.util import Pt, Emu, Cm
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn

# --- 定数 ---
TEMPLATE_PATH = "template/template.pptx"
DATA_PATH = "output/20260404/updates_ja.json"
OUTPUT_PATH = "output/20260404/azure-updates.pptx"

# テンプレートレイアウトインデックス（template 実測値）
LAYOUT_COVER = 0        # "1_Title Slide" (ph idx=0: Title, idx=12: Subtitle)
LAYOUT_CONTENT = 1      # "タイトルとコンテンツ" (ph idx=0: Title, idx=10: Content)
LAYOUT_SECTION = 3      # "Section Title" (ph idx=0: Title)

# カテゴリ表示順
CATEGORY_ORDER = [
    "セキュリティ",
    "ネットワーク",
    "コンピューティング",
    "データ & AI",
    "DevOps & 開発者ツール",
    "管理 & ガバナンス",
    "ハイブリッド & マルチクラウド",
    "IoT",
    "その他",
]

# 重要度ソート順
IMPORTANCE_ORDER = {"high": 0, "medium": 1, "low": 2}
IMPORTANCE_LABELS = {"high": "高", "medium": "中", "low": "低"}

# フォント設定
FONT_NAME = "Yu Gothic UI"

COLOR_TEXT = RGBColor(0x33, 0x33, 0x33)
COLOR_BLUE = RGBColor(0x00, 0x78, 0xD4)
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_GRAY = RGBColor(0x66, 0x66, 0x66)

# ステータスバッジ色
STATUS_COLORS = {
    "一般提供開始": RGBColor(0x10, 0x7C, 0x10),
    "パブリックプレビュー": RGBColor(0xCA, 0x50, 0x10),
    "廃止予定": RGBColor(0xD1, 0x34, 0x38),
}

# コンテンツプレースホルダーを拡張するジオメトリ (EMU)
CONTENT_LEFT = Emu(457200)      # 左位置 約 0.5 inch
CONTENT_TOP = Emu(939983)
CONTENT_HEIGHT = Emu(5460817)   # スライド下端近くまで拡張
CONTENT_WIDTH = Emu(10561638)   # 左オフセット分を差し引いた幅

# テキストフレーム内部余白 (EMU)
TF_MARGIN_LEFT = Emu(365760)    # 約 0.4 inch
TF_MARGIN_RIGHT = Emu(365760)
TF_MARGIN_TOP = Emu(91440)      # 約 0.1 inch
TF_MARGIN_BOTTOM = Emu(91440)

# 箇条書き上限
MAX_BULLETS = 5

# 箇条書きインデント (EMU) — ・の幅分ぶら下げ
BULLET_INDENT = 228600   # 約 0.25 inch = 0.64 cm

# 説明文1行あたり最大文字数（超えたら要約的に短縮）
MAX_DESC_CHARS = 120


# --- ヘルパー関数 ---

def _set_font(run, size, bold=False, color=None):
    """run のフォントプロパティを設定する（Yu Gothic UI / East Asian 込み）。"""
    run.font.name = FONT_NAME
    run.font.size = size
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    # East Asian typeface
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", FONT_NAME)


def _add_para(tf, text, size, bold=False, color=None, space_before=None, indent=None, alignment=None):
    """テキストフレームに段落を追加する。"""
    p = tf.add_paragraph()
    if space_before is not None:
        p.space_before = Pt(space_before)
    if indent is not None:
        p.level = 0
        p_elem = p._p
        pPr = p_elem.get_or_add_pPr()
        pPr.set("marL", str(indent))       # 左マージン (EMU)
        pPr.set("indent", str(-indent))     # ぶら下げインデント (first line offset)
    if alignment is not None:
        p.alignment = alignment
    run = p.add_run()
    run.text = text
    _set_font(run, size, bold=bold, color=color)
    return p


def _clear_text_frame(tf):
    """テキストフレームの既存テキストをクリアする。"""
    for i in range(len(tf.paragraphs) - 1, 0, -1):
        p_elem = tf.paragraphs[i]._p
        p_elem.getparent().remove(p_elem)
    tf.paragraphs[0].clear()


def _clean_bullet(text: str) -> str:
    """先頭の箇条書きマーカーを除去する。"""
    return re.sub(r"^[\s]*[•·・\-‐]\s*", "", text)


def _split_description(desc: str) -> list[str]:
    """説明文を文単位で分割し、箇条書きリストを返す。"""
    if not desc:
        return []
    desc = _clean_bullet(desc)
    # 「。」で分割
    parts = [p.strip() for p in re.split(r"。", desc) if p.strip()]
    result = []
    for p in parts[:MAX_BULLETS]:
        # 先頭のマーカーを除去（重複防止）
        p = re.sub(r"^[\s]*[•·・\-‐]\s*", "", p)
        # 長すぎる場合は末尾を省略
        if len(p) > MAX_DESC_CHARS:
            p = p[:MAX_DESC_CHARS] + "…"
        result.append(p)
    return result


def _format_availabilities(avails) -> str:
    if not avails:
        return ""
    parts = []
    for a in avails:
        ring = a.get("ring", "")
        year = a.get("year", "")
        month = a.get("month", "")
        parts.append(f"{ring} {year}/{month}" if year and month else ring)
    return ", ".join(parts)


def _delete_existing_slides(prs):
    """テンプレートの既存スライドを削除する。"""
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        rId = sldId.get(qn("r:id"))
        prs.part.drop_rel(rId)
        sldIdLst.remove(sldId)


def _resize_content_ph(ph):
    """コンテンツプレースホルダーをスライド下端まで拡張し、内部余白を設定する。"""
    ph.left = CONTENT_LEFT
    ph.top = CONTENT_TOP
    ph.height = CONTENT_HEIGHT
    ph.width = CONTENT_WIDTH
    tf = ph.text_frame
    tf.margin_left = TF_MARGIN_LEFT
    tf.margin_right = TF_MARGIN_RIGHT
    tf.margin_top = TF_MARGIN_TOP
    tf.margin_bottom = TF_MARGIN_BOTTOM


# --- データ処理 ---

def load_data():
    with open(DATA_PATH, "r", encoding="utf-8") as f:
        return json.load(f)


def group_and_sort(updates):
    """カテゴリ→重要度→日付でソート・グルーピング。"""
    grouped = OrderedDict()
    for cat in CATEGORY_ORDER:
        grouped[cat] = []

    for item in updates:
        cat = item.get("solutionCategory", "その他")
        if cat not in grouped:
            cat = "その他"
        grouped[cat].append(item)

    for cat in grouped:
        grouped[cat].sort(
            key=lambda x: (
                IMPORTANCE_ORDER.get(x.get("importance", "low"), 2),
                -(datetime.fromisoformat(
                    x["created"].replace("Z", "+00:00")).timestamp()
                  if x.get("created") else 0),
            )
        )

    return OrderedDict((k, v) for k, v in grouped.items() if v)


# --- スライド生成 ---

def create_cover_slide(prs):
    """表紙スライド。"""
    layout = prs.slide_layouts[LAYOUT_COVER]
    slide = prs.slides.add_slide(layout)

    title_ph = slide.placeholders[0]
    _clear_text_frame(title_ph.text_frame)
    run = title_ph.text_frame.paragraphs[0].add_run()
    run.text = "Azure Updates まとめ"
    _set_font(run, Pt(36), bold=True, color=COLOR_WHITE)

    sub_ph = slide.placeholders[12]
    _clear_text_frame(sub_ph.text_frame)
    run = sub_ph.text_frame.paragraphs[0].add_run()
    run.text = "2026/03/09 〜 2026/04/03"
    _set_font(run, Pt(20), color=COLOR_WHITE)


def create_summary_slide(prs, updates):
    """サマリースライド。"""
    layout = prs.slide_layouts[LAYOUT_CONTENT]
    slide = prs.slides.add_slide(layout)

    # タイトル
    title_ph = slide.placeholders[0]
    _clear_text_frame(title_ph.text_frame)
    run = title_ph.text_frame.paragraphs[0].add_run()
    run.text = "サマリー"
    _set_font(run, Pt(28), bold=True, color=COLOR_BLUE)

    # コンテンツ
    content_ph = slide.placeholders[10]
    _resize_content_ph(content_ph)
    tf = content_ph.text_frame
    tf.word_wrap = True
    _clear_text_frame(tf)

    # ステータス別件数
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "ステータス別件数"
    _set_font(run, Pt(22), bold=True, color=COLOR_BLUE)

    status_counts = Counter(item.get("status", "不明") for item in updates)
    for status in ["一般提供開始", "パブリックプレビュー", "廃止予定"]:
        count = status_counts.pop(status, 0)
        if count > 0:
            badge_color = STATUS_COLORS.get(status, COLOR_TEXT)
            p = _add_para(tf, f"{status}: {count}件", Pt(18), color=badge_color, space_before=2)
            p.level = 1
    for status, count in status_counts.items():
        if count > 0:
            p = _add_para(tf, f"{status}: {count}件", Pt(18), color=COLOR_TEXT, space_before=2)
            p.level = 1

    # カテゴリ別件数
    _add_para(tf, "ソリューションカテゴリ別件数", Pt(22), bold=True, color=COLOR_BLUE, space_before=16)

    cat_counts = Counter(item.get("solutionCategory", "その他") for item in updates)
    for cat in CATEGORY_ORDER:
        count = cat_counts.get(cat, 0)
        if count > 0:
            p = _add_para(tf, f"{cat}: {count}件", Pt(18), color=COLOR_TEXT, space_before=2)
            p.level = 1


def create_section_slide(prs, category_name, count):
    """カテゴリセクション区切りスライド。"""
    layout = prs.slide_layouts[LAYOUT_SECTION]
    slide = prs.slides.add_slide(layout)

    title_ph = slide.placeholders[0]
    _clear_text_frame(title_ph.text_frame)
    run = title_ph.text_frame.paragraphs[0].add_run()
    run.text = f"{category_name}（{count}件）"
    _set_font(run, Pt(36), bold=True, color=COLOR_WHITE)


def create_update_slide(prs, item):
    """個別更新スライド。"""
    layout = prs.slide_layouts[LAYOUT_CONTENT]
    slide = prs.slides.add_slide(layout)

    # --- タイトル ---
    title_ph = slide.placeholders[0]
    _clear_text_frame(title_ph.text_frame)
    title_text = item.get("title_ja", item.get("title_en", ""))

    # タイトルの自動縮小を有効化
    title_ph.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    title_ph.text_frame.word_wrap = True

    run = title_ph.text_frame.paragraphs[0].add_run()
    run.text = title_text
    title_size = Pt(20) if len(title_text) > 60 else Pt(24)
    _set_font(run, title_size, bold=True, color=COLOR_BLUE)

    # --- コンテンツ ---
    content_ph = slide.placeholders[10]
    _resize_content_ph(content_ph)
    tf = content_ph.text_frame
    tf.word_wrap = True
    _clear_text_frame(tf)

    # ステータス + 重要度バッジ （1行目）
    status = item.get("status", "不明")
    importance = item.get("importance", "low")
    badge_color = STATUS_COLORS.get(status, COLOR_TEXT)
    imp_label = IMPORTANCE_LABELS.get(importance, importance)

    p = tf.paragraphs[0]
    p.space_after = Pt(8)
    r1 = p.add_run()
    r1.text = f"【{status}】"
    _set_font(r1, Pt(16), bold=True, color=badge_color)
    r2 = p.add_run()
    r2.text = f"  重要度: {imp_label}"
    _set_font(r2, Pt(16), bold=True, color=COLOR_TEXT)

    # 説明（箇条書き — テンプレートのビルトイン箇条書きマーカーを使用）
    for line in _split_description(item.get("description_ja", "")):
        p = _add_para(tf, line, Pt(14), color=COLOR_TEXT, space_before=3)
        # テンプレート側の箇条書きマーカーを使うため level=1 を設定
        p.level = 1

    # --- 付帯情報（区切り線代わりに少し広めのスペース）---
    # 対象製品
    products = item.get("products", [])
    if products:
        _add_para(tf, f"対象製品: {', '.join(products)}", Pt(13),
                  bold=True, color=COLOR_GRAY, space_before=14)

    # 作成日
    created = item.get("created", "")
    if created:
        try:
            dt = datetime.fromisoformat(created.replace("Z", "+00:00"))
            date_str = dt.strftime("%Y/%m/%d")
        except ValueError:
            date_str = created[:10]
        _add_para(tf, f"作成日: {date_str}", Pt(13), color=COLOR_GRAY, space_before=2)

    # 可用性
    avail_text = _format_availabilities(item.get("availabilities", []))
    if avail_text:
        _add_para(tf, f"可用性: {avail_text}", Pt(13), color=COLOR_GRAY, space_before=2)

    # 詳細リンク
    link_url = item.get("link", "")
    if not link_url:
        item_id = item.get("id", "")
        if item_id:
            link_url = f"https://azure.microsoft.com/updates/?id={item_id}"
    if link_url:
        p = tf.add_paragraph()
        p.space_before = Pt(2)
        run_label = p.add_run()
        run_label.text = "詳細: "
        _set_font(run_label, Pt(13), color=COLOR_GRAY)
        run_link = p.add_run()
        run_link.text = link_url
        _set_font(run_link, Pt(13), color=COLOR_BLUE)
        run_link.hyperlink.address = link_url


# --- メイン処理 ---

def main():
    print(f"Loading data from {DATA_PATH} ...")
    updates = load_data()
    print(f"  {len(updates)} items loaded.")

    print(f"Loading template from {TEMPLATE_PATH} ...")
    prs = Presentation(TEMPLATE_PATH)

    print("Deleting existing template slides ...")
    _delete_existing_slides(prs)

    print("Creating cover slide ...")
    create_cover_slide(prs)

    print("Creating summary slide ...")
    create_summary_slide(prs, updates)

    print("Grouping and sorting updates ...")
    grouped = group_and_sort(updates)

    slide_count = 2  # cover + summary
    for category, items in grouped.items():
        print(f"  {category}: {len(items)} items")
        create_section_slide(prs, category, len(items))
        slide_count += 1
        for item in items:
            create_update_slide(prs, item)
            slide_count += 1

    print(f"Total slides: {slide_count}")
    print(f"Saving to {OUTPUT_PATH} ...")
    prs.save(OUTPUT_PATH)
    print("Done!")


if __name__ == "__main__":
    main()
